import json, datetime, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from clusters import C
from opps import OPPS, SEARCHES

S=__import__('os').path.join(__import__('os').path.dirname(__file__),'..','data')
R=json.load(open(S+'/kanallar_ham.json'))
code_of={n:k for k,(_,l) in C.items() for n in l}
TODAY=datetime.date(2026,10,9)
F=Font(name='Arial',size=10); FB=Font(name='Arial',size=10,bold=True); FH=Font(name='Arial',size=10,bold=True,color='FFFFFF')
HF=PatternFill('solid',fgColor='1F3864'); WR=Alignment(wrap_text=True,vertical='top')
def header(ws,cols,row=1):
    for i,(h,w) in enumerate(cols,1):
        c=ws.cell(row=row,column=i,value=h); c.font=FH; c.fill=HF; c.alignment=Alignment(wrap_text=True,vertical='center')
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes=ws.cell(row=row+1,column=3)

wb=Workbook()
# ---------- Özet
ws=wb.active; ws.title='Özet'
lines=[
 ("Konsept Araştırması – Nexlev verisi, 9 Ekim 2026",FB),
 ("",F),
 ("KRİTERLER",FB),
 ("• Kanal açılışı ≤ 5 ay (2026-05-09 sonrası) – Nexlev 'channelCreatedAfter' + ilk video tarihiyle çapraz kontrol edildi.",F),
 ("• Para kazanma açık (isMonetizationEnabled = true).",F),
 ("• Tahmini aylık AdSense geliri ≥ $1.000 (Nexlev tahmini; sponsorluk/ürün dahil değil).",F),
 ("",F),
 ("KAPSAM VE YÖNTEM",FB),
 ("1) search_niche_finder_channels (query='*', yukarıdaki filtreler): Nexlev 1.137 kanal buldu. Ücretsiz plan 24 saatte 5 sorgu izin verdiği için 5 sayfa = 500 kanal çekilebildi (sayfa 1,2,3,4,6).",F),
 ("2) latest_discovered_faceless_niches (1.000 kanal): yaşı ≤155 gün, monetize, ≥$1.000 olan 71 kanal; 61'i yeni → toplam 561 kanal.",F),
 ("3) 561 kanalın her birinin son video başlıkları tek tek okunup 101 konsept kümesine elle atandı (Kanallar sayfası 'Konsept' sütunu).",F),
 ("4) Aday boşluklar için 13 canlı YouTube araması yapıldı (rakip sayısı + kopyaların izlenmesi) ve 1.310 güncel faceless outlier video tarandı → 'Rakip Kontrolleri' sayfası.",F),
 ("5) 'Konseptler' sayfasında her küme için yeni kanal sayısı/geliri formülle hesaplanır; 'Fırsatlar' sayfası puanlı öneri listesidir.",F),
 ("",F),
 ("ÖNEMLİ UYARILAR",FB),
 ("• 1.137 kanalın 637'si henüz çekilemedi (kota). Yarın kalan sayfalar (5, 7-12) çekilip tablo güncellenebilir.",F),
 ("• Gelir rakamları Nexlev tahminidir. 1-3K aboneli kanalda $10K+/ay gibi değerler RPM tahmininden şişmiş olabilir – kendi kontrolünde izlenme/RPM'ye bak.",F),
 ("• WIRED Investigations (#351) 'yeni kanal' görünse de ilk videosu 2016 – kurumsal yeniden yükleme, hariç tut.",F),
 ("• 'Latest' kaynaklı 61 kanalda kanal yaşı gün cinsinden (Nexlev ay yuvarlaması: 14/21/28/30/61/92/122/153).",F),
 ("• Şüpheli gelir tahminleri: 3DVerse 12 (#2, $62.8K, video başı ~4K izlenme), 3D_Spectra_YT (#3), Senju (#1), Crazy Football (#6), Baseplate (#16), REDDIT HAN (#42) – izlenmeleri gelirle orantısız; muhtemelen Nexlev hatası veya Shorts ağırlıklı. PopDaBalloon (#4) TV şovu yeniden yüklemesi (telif riski).",F),
 ("• Veri setindeki kanalların tamamı faceless (Nexlev veritabanı faceless odaklı). %35'i Nexlev'e göre AI içerik.",F),
 ("",F),
 ("ANA BULGU – 'boşluk' nasıl okunmalı",FB),
 ("Başarılı bir format 1-3 hafta içinde 10-15 birebir kopya kanal çekiyor (Amish, Haile Selassie çocukları, 'cop caught lying in court'). Kopyalar 0-2K izlenmede ölüyor; orijinal kazanmaya devam ediyor.",F),
 ("Yani gerçek boşluk = (a) talep kanıtlı ama faceless arzı olmayan konu, ya da (b) kanıtlı ŞABLONUN henüz kimsenin girmediği komşu dikeye taşınması. Fırsatlar sayfası bu mantıkla sıralandı.",F),
 ("",F),
 ("SAYFALAR",FB),
 ("Kanallar – 561 kanal, link, metrik, konsept, son videolar (filtrelenebilir tablo)",F),
 ("Konseptler – 101 küme: yeni kanal sayısı, toplam/ortalama/maks gelir (formül), doygunluk, dış kontrol notu",F),
 ("Fırsatlar – puanlı boşluk/konsept önerileri, kanıt ve önerilen açı",F),
 ("Rakip Kontrolleri – yapılan YouTube aramaları ve sonuçları",F),
]
for i,(t,f) in enumerate(lines,1):
    c=ws.cell(row=i,column=1,value=t); c.font=f; c.alignment=Alignment(wrap_text=True,vertical='top')
ws['A1'].font=Font(name='Arial',size=14,bold=True)
ws.column_dimensions['A'].width=150

# ---------- Kanallar
wk=wb.create_sheet('Kanallar')
cols=[("No",6),("Kanal",28),("Link",14),("Konsept kodu",14),("Konsept",34),("Dil",6),("Ülke",6),("Nexlev kategori",20),("Format",13),
 ("Aylık gelir ($)",12),("Abone",10),("Video sayısı",9),("Ort. izlenme/video",12),("Aylık izlenme",12),("Ort. süre (dk)",9),("RPM ($)",8),
 ("Outlier",8),("AI içerik",8),("İlk video / yaş",13),("Kanal yaşı (gün)",10),("Kaynak",8),("Son videolar (izlenme)",90)]
header(wk,cols)
for i,r in enumerate(R,2):
    code=code_of[r['n']]; first=r['first'] or ''
    age = (TODAY-datetime.date.fromisoformat(first)).days if first else r.get('days')
    vals=[r['n'],r['title'],r['url'],code,C[code][0],r['lang'] or '',r['loc'] or '',r['cat'] or ', '.join(r['tags'][:3]),r['fmt'] or '',
          r['rev'],r['subs'],r['vids'],r['avgv'],r['mviews'],r['len'],round(r['rpm'],2) if r['rpm'] else None,r['outlier'],
          'Evet' if r['ai'] else 'Hayır',first or f"{r.get('days')} gün",age,r['src'],
          ' | '.join(f"{t} ({v:,})" for t,v,_ in r['vt'][:4])]
    for j,v in enumerate(vals,1):
        c=wk.cell(row=i,column=j,value=v); c.font=F
    lc=wk.cell(row=i,column=3); lc.hyperlink=r['url']; lc.value='YouTube'; lc.font=Font(name='Arial',size=10,color='0563C1',underline='single')
    for j,fmt in [(10,'$#,##0'),(11,'#,##0'),(12,'#,##0'),(13,'#,##0'),(14,'#,##0'),(16,'0.00')]:
        wk.cell(row=i,column=j).number_format=fmt
n=len(R)+1
t=Table(displayName='Kanallar',ref=f"A1:{get_column_letter(len(cols))}{n}"); t.tableStyleInfo=TableStyleInfo(name='TableStyleLight9',showRowStripes=True); wk.add_table(t)

# ---------- Konseptler
CHECK={
 "EBIKEPOV":"YouTube: tek ciddi kanal (Yucky Stew) → AÇIK","SHERLOCK":"YouTube: 4-5 küçük kanal → KISMEN AÇIK",
 "AMISH":"YouTube: 10+ kopya, kopyalar 0-2K → KOPYA DOLU","BATTLEGRAVE":"YouTube: kopyalar 4-36K alıyor → KISMEN AÇIK",
 "AFTERCROWN":"YouTube: ~15 kopya <3K → KOPYA DOLU","COPCOURT":"YouTube: 15+ kopya 0-500 → DOLU",
 "USPOVERTY":"YouTube: talep 10M+, faceless arz yok → AÇIK","SENIORDIY":"YouTube: ~6 kanal hepsi izlenme alıyor → KISMEN DOLU",
 "OCEAN4K":"YouTube: onlarca kanal 12-2K → DOLU","NOBODYVISIT":"YouTube: ~8 kanal 20-100K → KISMEN AÇIK",
 "FINANCE":"'Economics of owning' YouTube: 15+ kanal → DOLU","BIZFALL":"Golf dikeyi: tek lider (Persimmon) → DİKEYDE AÇIK",
 "JAPANLIFE":"YouTube: 15+ kanal ~100 izlenme → DOLU","CLASSICTV":"Outlier: Curtain Cut (dizi-içi trivia) → AÇIK açı",
 "MILHIST":"Outlier: Kore Savaşı kanalı 286K → WW2 dışı AÇIK",
}
wc=wb.create_sheet('Konseptler')
ccols=[("Kod",14),("Konsept",50),("Yeni kanal sayısı",11),("Toplam aylık gelir ($)",14),("Ort. gelir/kanal ($)",13),("En yüksek kanal geliri ($)",14),
       ("Doygunluk (yeni kanal sayısına göre)",20),("Dış kontrol (YouTube arama)",48),("Kanallar (No)",40)]
header(wc,ccols)
items=sorted(C.items(),key=lambda kv:-sum(next(r['rev'] for r in R if r['n']==n) for n in kv[1][1]))
for i,(k,(name,ns)) in enumerate(items,2):
    rng=f"Kanallar!$D$2:$D${n}"
    vals=[k,name,f'=COUNTIF({rng},A{i})',f'=SUMIF({rng},A{i},Kanallar!$J$2:$J${n})',f'=IF(C{i}=0,0,D{i}/C{i})',
          f'=_xlfn.MAXIFS(Kanallar!$J$2:$J${n},{rng},A{i})',
          f'=IF(C{i}>=6,"Dolu (6+)",IF(C{i}>=3,"Kısmen dolu (3-5)","Az rakip (1-2)"))',CHECK.get(k,'—'),', '.join(f'#{x}' for x in sorted(ns))]
    for j,v in enumerate(vals,1):
        c=wc.cell(row=i,column=j,value=v); c.font=F
        if j in (2,8,9): c.alignment=WR
    for j in (4,5,6): wc.cell(row=i,column=j).number_format='$#,##0'
last=len(items)+1
wc.cell(row=last+1,column=2,value='TOPLAM').font=FB
wc.cell(row=last+1,column=3,value=f'=SUM(C2:C{last})').font=FB
c=wc.cell(row=last+1,column=4,value=f'=SUM(D2:D{last})'); c.font=FB; c.number_format='$#,##0'
wc.cell(row=last+2,column=2,value='Doygunluk eşiği: aynı kümede son 5 ayda açılmış ve ≥$1K/ay kazanan kanal sayısı (yalnız bu veri setinde). Dış kontrol sütunu canlı YouTube aramasından elle yazıldı.').font=Font(name='Arial',size=9,italic=True)
t=Table(displayName='Konseptler',ref=f"A1:I{last}"); t.tableStyleInfo=TableStyleInfo(name='TableStyleLight9',showRowStripes=True); wc.add_table(t)

# ---------- Fırsatlar
wf=wb.create_sheet('Fırsatlar')
fcols=[("Sıra",6),("Konsept / fırsat",38),("Tip",20),("Kanıt (kanal no = Kanallar sayfası)",50),("Rakip kontrolü",42),("Önerilen açı / içerik planı",60),("Risk",30),("Puan (1-10)",9)]
header(wf,fcols)
for i,o in enumerate(sorted(OPPS,key=lambda o:-o[6]),2):
    for j,v in enumerate([i-1,*o],1):
        c=wf.cell(row=i,column=j,value=v); c.font=F; c.alignment=WR
    wf.cell(row=i,column=8).fill=PatternFill('solid',fgColor='FFFF00')
r0=len(OPPS)+3
wf.cell(row=r0,column=2,value='Puan = talep kanıtı (kanal geliri/outlier) + rakip azlığı + faceless/AI üretilebilirlik + RPM tahmini; benim değerlendirmem (sarı hücreler değiştirilebilir).').font=Font(name='Arial',size=9,italic=True)

# ---------- Rakip Kontrolleri
wr=wb.create_sheet('Rakip Kontrolleri')
rcols=[("Arama sorgusu",40),("Tarih filtresi",12),("Bulgular",90),("Sonuç",36)]
header(wr,rcols)
for i,s in enumerate(SEARCHES,2):
    for j,v in enumerate(s,1):
        c=wr.cell(row=i,column=j,value=v); c.font=F; c.alignment=WR
out=sys.argv[1]; wb.save(out); print('saved',out)
